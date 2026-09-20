"""Voice transcription, image OCR/vision, and document text extraction — Sarvam AI only.

All AI processing goes through Sarvam AI (per user directive). No OpenAI / Anthropic / Gemini.

- Audio transcription: Sarvam Speech-to-Text (saaras:v3, auto-detect Hindi/English/Hinglish).
- Image OCR/vision: Sarvam Vision (Indic-first VLM).
- Document text: local extraction with pypdf / python-docx / openpyxl.
"""
import os
import io
import base64
import logging
import tempfile
from pathlib import Path

import httpx

logger = logging.getLogger(__name__)

SARVAM_API_KEY = os.environ["SARVAM_API_KEY"]
SARVAM_STT_URL = "https://api.sarvam.ai/speech-to-text"
SARVAM_VISION_URL = "https://api.sarvam.ai/v1/vision"

# Sarvam BCP-47 codes we accept from clients (everything else -> auto-detect).
_LANG_MAP = {
    "en": "en-IN", "en-in": "en-IN", "english": "en-IN",
    "hi": "hi-IN", "hi-in": "hi-IN", "hindi": "hi-IN",
    "bn": "bn-IN", "gu": "gu-IN", "kn": "kn-IN", "ml": "ml-IN",
    "mr": "mr-IN", "od": "od-IN", "pa": "pa-IN", "ta": "ta-IN",
    "te": "te-IN",
}


def _resolve_lang(language: str | None) -> str:
    """Return a Sarvam-compatible BCP-47 code, or 'unknown' for Hinglish/auto."""
    if not language:
        return "unknown"
    key = language.strip().lower()
    if key in ("auto", "unknown", "hinglish"):
        return "unknown"
    return _LANG_MAP.get(key, "unknown")


async def transcribe_audio(audio_bytes: bytes, filename: str, language: str = "auto") -> str:
    """Transcribe an audio clip via Sarvam STT. Auto-detects Hindi/English/Hinglish."""
    suffix = Path(filename).suffix.lower() or ".m4a"
    # Normalise to something Sarvam accepts.
    if suffix not in {".m4a", ".mp3", ".wav", ".webm", ".mp4", ".mpeg", ".mpga", ".aac", ".ogg", ".flac", ".opus", ".amr"}:
        suffix = ".m4a"
    tmp_name = None
    try:
        with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
            tmp.write(audio_bytes)
            tmp_name = tmp.name

        lang = _resolve_lang(language)
        headers = {"api-subscription-key": SARVAM_API_KEY}

        with open(tmp_name, "rb") as fh:
            files = {"file": (Path(filename or "audio").name, fh, "application/octet-stream")}
            data = {"model": "saaras:v3", "language_code": lang, "with_timestamps": "false"}
            async with httpx.AsyncClient(timeout=90) as client:
                resp = await client.post(SARVAM_STT_URL, headers=headers, data=data, files=files)

        if resp.status_code >= 400:
            logger.warning("Sarvam STT %s: %s", resp.status_code, resp.text[:200])
            resp.raise_for_status()
        payload = resp.json() if resp.content else {}
        text = payload.get("transcript") or payload.get("text") or ""
        return (text or "").strip()
    finally:
        if tmp_name:
            Path(tmp_name).unlink(missing_ok=True)


async def image_qa(image_bytes: bytes, mime_type: str, question: str) -> str:
    """OCR / vision Q&A via Sarvam Vision. Understands English, Hindi and Hinglish."""
    image_b64 = base64.b64encode(image_bytes).decode("ascii")
    body = {
        "image": f"data:{mime_type or 'image/jpeg'};base64,{image_b64}",
        "prompt": (
            "Analyze this user-provided image (screenshot, invoice, photo). Read all visible text "
            "with OCR. Preserve numbers, dates and amounts exactly; never invent missing values. "
            "You understand English, Hindi and Hinglish. Answer the user's question below.\n\n"
            f"Question: {question}"
        ),
    }
    # Sarvam Vision accepts either the bearer or api-subscription-key header depending on
    # the account. Try the subscription-key header first (same as STT), then fall back.
    headers_variants = [
        {"api-subscription-key": SARVAM_API_KEY, "Content-Type": "application/json"},
        {"Authorization": f"Bearer {SARVAM_API_KEY}", "Content-Type": "application/json"},
    ]
    last: httpx.Response | None = None
    async with httpx.AsyncClient(timeout=90) as client:
        for h in headers_variants:
            resp = await client.post(SARVAM_VISION_URL, headers=h, json=body)
            last = resp
            if resp.status_code < 400:
                data = resp.json()
                for key in ("answer", "response", "text", "output"):
                    if isinstance(data.get(key), str):
                        return data[key].strip()
                # Some responses may nest content under choices.
                choices = data.get("choices")
                if isinstance(choices, list) and choices:
                    msg = choices[0].get("message") or {}
                    if isinstance(msg.get("content"), str):
                        return msg["content"].strip()
                return str(data)[:2000]
    if last is not None:
        logger.warning("Sarvam Vision %s: %s", last.status_code, last.text[:200])
        last.raise_for_status()
    return ""


def extract_document_text(data: bytes, filename: str, mime: str) -> str:
    """Best-effort text extraction from common document formats. Local only — no LLM."""
    ext = Path(filename).suffix.lower()
    try:
        if ext == ".pdf" or mime == "application/pdf":
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(data))
            return "\n".join((p.extract_text() or "") for p in reader.pages[:40]).strip()
        if ext in (".docx",):
            import docx
            d = docx.Document(io.BytesIO(data))
            return "\n".join(p.text for p in d.paragraphs).strip()
        if ext in (".xlsx",):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets[:5]:
                out.append(f"# Sheet: {ws.title}")
                for row in ws.iter_rows(values_only=True):
                    cells = [str(c) for c in row if c is not None]
                    if cells:
                        out.append(" | ".join(cells))
            return "\n".join(out[:1000]).strip()
        if ext in (".csv", ".txt", ".md", ".json"):
            return data.decode("utf-8", errors="ignore")[:60000].strip()
    except Exception as e:
        logger.warning(f"Text extraction failed for {filename}: {e}")
    return ""
